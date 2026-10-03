from contextlib import asynccontextmanager

from fastapi import FastAPI, Depends, HTTPException, Body
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy.orm import Session
from typing import List, Dict, Any

from database import engine, get_db
import models  # noqa: F401  (registers ORM tables)
import profile_manager
import tracker
from field_matcher import match_fields
from api_errors import RequestIDMiddleware, register_error_handlers
from identity import CurrentUser, get_current_user
from schema_guard import assert_schema_current
from app_tracker.router import router as applications_router


@asynccontextmanager
async def lifespan(_app: FastAPI):
    # Schema is managed by Alembic migrations (`cd backend && alembic upgrade head`).
    # Refuse to serve requests against an outdated database instead of failing per request.
    assert_schema_current(engine)
    yield


app = FastAPI(title="Job Copilot API", lifespan=lifespan)

# Allow CORS for extension
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "chrome-extension://mhbjkfhgfflooeplobpgpicndphcminm",
        "http://localhost:3000"
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.add_middleware(RequestIDMiddleware)
register_error_handlers(app)
app.include_router(applications_router)

# API Endpoints
@app.post("/api/match-fields")
def api_match_fields(fields: List[Dict[str, Any]] = Body(...), db: Session = Depends(get_db)):
    profile = profile_manager.get_profile(db)
    matches = match_fields(fields, profile)
    return matches

@app.post("/api/log-application")
def api_log_application(data: Dict[str, Any] = Body(...), db: Session = Depends(get_db),
                        user: CurrentUser = Depends(get_current_user)):
    application = tracker.log_application(
        db,
        user.id,
        url=str(data.get("url", ""))[:2048],
        title=str(data.get("title", ""))[:300],
        company=str(data.get("company", ""))[:300],
        platform=str(data.get("platform", "other"))[:50],
        job_description=str(data.get("job_description", ""))[:50_000]
    )
    return {"status": "success", "id": application.id}

@app.get("/api/profile")
def api_get_profile(db: Session = Depends(get_db)):
    return profile_manager.get_profile(db)

@app.put("/api/profile")
def api_update_profile(data: Dict[str, Any] = Body(...), db: Session = Depends(get_db)):
    return profile_manager.update_profile(db, data)

@app.get("/api/stats")
def api_get_stats(db: Session = Depends(get_db), user: CurrentUser = Depends(get_current_user)):
    return tracker.get_stats(db, user.id)

@app.get("/api/applications")
def api_get_applications(limit: int = 100, offset: int = 0, db: Session = Depends(get_db),
                         user: CurrentUser = Depends(get_current_user)):
    return tracker.get_applications(db, user.id, limit, offset)

@app.put("/api/applications/{app_id}")
def api_update_application(app_id: int, data: Dict[str, str] = Body(...), db: Session = Depends(get_db),
                           user: CurrentUser = Depends(get_current_user)):
    status = data.get("status")
    notes = data.get("notes")
    if not status and notes is None:
        raise HTTPException(status_code=400, detail="status or notes required")
    # Not found -> 404, unknown status -> 400, disallowed transition -> 409 (all as {"detail": ...}).
    tracker.update_application(db, user.id, app_id, status=status or None, notes=notes)
    return {"status": "success"}

@app.delete("/api/applications/{app_id}")
def api_delete_application(app_id: int, db: Session = Depends(get_db),
                           user: CurrentUser = Depends(get_current_user)):
    tracker.delete_application(db, user.id, app_id)
    return {"status": "success"}

import io
@app.get("/api/export")
def api_export_applications(db: Session = Depends(get_db), user: CurrentUser = Depends(get_current_user)):
    csv_data = tracker.export_applications_csv(db, user.id)
    return StreamingResponse(
        io.StringIO(csv_data),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=applications.csv"}
    )

# --- AI Writer Endpoints ---
import ai_writer

@app.post("/api/generate-cover-letter")
def api_generate_cover_letter(data: Dict[str, Any] = Body(...), db: Session = Depends(get_db)):
    job_description = data.get("job_description", "")
    company_name = data.get("company_name", "")
    target_product = data.get("target_product", "")
    
    if not job_description:
        raise HTTPException(status_code=400, detail="job_description required")
    profile = profile_manager.get_profile(db)
    text = ai_writer.generate_cover_letter(job_description, profile, company_name, target_product)
    return {"cover_letter": text}

@app.post("/api/generate-answer")
def api_generate_answer(data: Dict[str, Any] = Body(...), db: Session = Depends(get_db)):
    question = data.get("question", "")
    job_description = data.get("job_description", "")
    if not question:
        raise HTTPException(status_code=400, detail="question required")
    profile = profile_manager.get_profile(db)
    text = ai_writer.generate_answer(question, profile, job_description)
    return {"answer": text}

@app.post("/api/tailor-resume")
def api_tailor_resume(data: Dict[str, Any] = Body(...), db: Session = Depends(get_db)):
    bullets = data.get("bullets", [])
    job_description = data.get("job_description", "")
    if not bullets or not job_description:
        raise HTTPException(status_code=400, detail="bullets and job_description required")
    profile = profile_manager.get_profile(db)
    result = ai_writer.tailor_resume_bullets(bullets, job_description, profile)
    return {"bullets": result}

@app.post("/api/humanize")
def api_humanize(data: Dict[str, str] = Body(...)):
    text = data.get("text", "")
    if not text:
        raise HTTPException(status_code=400, detail="text required")
    result = ai_writer.humanize_text(text)
    return {"text": result}

@app.post("/api/interview-prep")
def api_interview_prep(data: Dict[str, Any] = Body(...), db: Session = Depends(get_db)):
    job_description = data.get("job_description", "")
    if not job_description:
        raise HTTPException(status_code=400, detail="job_description required")
    profile = profile_manager.get_profile(db)
    result = ai_writer.generate_interview_prep(job_description, profile)
    return {"prep": result}

# --- Export Endpoints ---

from datetime import datetime

def number_to_words(n):
    words = {
        0: 'Zero', 1: 'One', 2: 'Two', 3: 'Three', 4: 'Four', 5: 'Five', 
        6: 'Six', 7: 'Seven', 8: 'Eight', 9: 'Nine', 10: 'Ten',
        11: 'Eleven', 12: 'Twelve', 13: 'Thirteen', 14: 'Fourteen', 
        15: 'Fifteen', 16: 'Sixteen', 17: 'Seventeen', 18: 'Eighteen', 19: 'Nineteen',
        20: 'Twenty', 30: 'Thirty', 40: 'Forty', 50: 'Fifty'
    }
    if n in words:
        return words[n]
    if n < 60:
        return words[(n // 10) * 10] + '_' + words[n % 10]
    return str(n)

def get_word_datetime():
    now = datetime.now()
    months = ['January', 'February', 'March', 'April', 'May', 'June', 
              'July', 'August', 'September', 'October', 'November', 'December']
    month = months[now.month - 1]
    
    ordinals = {
        1: 'First', 2: 'Second', 3: 'Third', 4: 'Fourth', 5: 'Fifth',
        6: 'Sixth', 7: 'Seventh', 8: 'Eighth', 9: 'Ninth', 10: 'Tenth',
        11: 'Eleventh', 12: 'Twelfth', 13: 'Thirteenth', 14: 'Fourteenth', 15: 'Fifteenth',
        16: 'Sixteenth', 17: 'Seventeenth', 18: 'Eighteenth', 19: 'Nineteenth', 20: 'Twentieth',
        21: 'Twenty_First', 22: 'Twenty_Second', 23: 'Twenty_Third', 24: 'Twenty_Fourth', 25: 'Twenty_Fifth',
        26: 'Twenty_Sixth', 27: 'Twenty_Seventh', 28: 'Twenty_Eighth', 29: 'Twenty_Ninth', 30: 'Thirtieth',
        31: 'Thirty_First'
    }
    day = ordinals.get(now.day, number_to_words(now.day))
    
    year_str = "Two_Thousand"
    if now.year > 2000:
        rem = now.year % 100
        if rem > 0:
            year_str = "Two_Thousand_" + number_to_words(rem)
        
    hour = now.hour % 12
    if hour == 0: hour = 12
    hour_str = number_to_words(hour)
    
    minute_str = number_to_words(now.minute) if now.minute > 0 else ""
    ampm = "AM" if now.hour < 12 else "PM"
    
    time_str = f"{hour_str}_{minute_str}_{ampm}" if minute_str else f"{hour_str}_{ampm}"
    
    return f"{month}_{day}_{year_str}_{time_str}"

from docx import Document
from fpdf import FPDF
import io


@app.post("/api/export-doc")
def api_export_document(data: Dict[str, Any] = Body(...), db: Session = Depends(get_db)):
    text = data.get("text", "")
    fmt = data.get("format", "tex")
    doc_type = data.get("doc_type", "Document")
    
    # 1. Fix Unicode Characters (smart quotes, dashes) that break basic fonts
    text = text.replace('’', "'").replace('‘', "'")
    text = text.replace('”', '"').replace('“', '"')
    text = text.replace('–', '-').replace('—', '--')
    text = text.replace('…', '...')
    text = text.replace('•', '-')
    text = text.replace("\'", "'")
    
    profile = profile_manager.get_profile(db)
    name = profile.get("name", "Applicant")
    first_name = name.split(" ")[0] if name else "Applicant"
    
    # Generate filename
    dt_str = get_word_datetime()
    safe_type = doc_type.replace(" ", "_")
    filename = f"{first_name}_{safe_type}_{dt_str}.{fmt}"
    
    # 2. Add Professional Header if it's a Cover Letter
    if doc_type == "Cover Letter":
        profile = profile_manager.get_profile(db)
        name = profile.get("name", "Applicant")
        phone = profile.get("phone", "")
        email = profile.get("email", "")
        location = profile.get("location", "")
        linkedin = profile.get("linkedin_url", "")
        
        contact_parts = [p for p in [phone, email, location] if p]
        contact_line = " | ".join(contact_parts)
        
        header = f"{name.upper()}\n"
        if contact_line:
            header += f"{contact_line}\n"
        if linkedin:
            header += f"{linkedin}\n"
        
        header += "_" * 60 + "\n\n\n"
        text = header + text

    if fmt == "tex":
        tex_content = "\\documentclass{article}\n\\begin{document}\n" 
        for paragraph in text.split('\n'):
            if paragraph.strip():
                p = paragraph.replace('\\', '\\\\').replace('%', '\\%').replace('$', '\\$').replace('_', '\\_').replace('#', '\\#')
                tex_content += p + "\n\n"
        tex_content += "\\end{document}"
        
        return StreamingResponse(
            io.StringIO(tex_content),
            media_type="application/x-tex",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    elif fmt == "docx":
        doc = Document()
        for paragraph in text.split('\n'):
            if paragraph.strip():
                if paragraph.startswith('_'*10):
                    pass # Skip the hardcoded line for docx, can use borders but skipping is easier
                else:
                    doc.add_paragraph(paragraph.strip())
        
        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)
        return StreamingResponse(
            bio,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    elif fmt == "pdf":
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("helvetica", size=11)
        
        # Clean to latin-1 for FPDF basic fonts
        cleaned_text = text.encode('latin-1', 'replace').decode('latin-1')
        
        for paragraph in cleaned_text.split('\n'):
            if paragraph.strip():
                if paragraph.startswith('_'*10):
                    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
                    pdf.ln(5)
                elif paragraph.isupper() and doc_type == "Cover Letter" and paragraph == cleaned_text.split('\n')[0]:
                    # Format Name in Header
                    pdf.set_font("helvetica", "B", 16)
                    pdf.cell(0, 8, paragraph.strip(), align='C', ln=True)
                    pdf.set_font("helvetica", size=10)
                elif "|" in paragraph and doc_type == "Cover Letter":
                    # Format Contact Info
                    pdf.cell(0, 5, paragraph.strip(), align='C', ln=True)
                elif "linkedin.com" in paragraph and doc_type == "Cover Letter":
                    pdf.cell(0, 5, paragraph.strip(), align='C', ln=True)
                else:
                    pdf.set_font("helvetica", size=11)
                    pdf.multi_cell(0, 6, paragraph.strip(), markdown=True)
                    pdf.ln(3)
        
        pdf_bytes = pdf.output()
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    
    raise HTTPException(status_code=400, detail="Invalid format")


# Mount static files for Dashboard
import os
static_dir = os.path.join(os.path.dirname(__file__), "static")
if os.path.exists(static_dir):
    app.mount("/static", StaticFiles(directory=static_dir), name="static")

@app.get("/")
def serve_dashboard():
    index_path = os.path.join(static_dir, "index.html")
    if os.path.exists(index_path):
        return FileResponse(index_path)
    return {"message": "Dashboard not found"}

